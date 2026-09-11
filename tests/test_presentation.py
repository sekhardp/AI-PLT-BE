import base64
import io
from fastapi.testclient import TestClient
from pptx import Presentation

from app.api.v1.schemas.presentation import (
    ChartConfig,
    ChartSeries,
    KPICard,
    Slide,
    SlideDeck,
)
from app.main import app
from app.services.presentation import PPTXBuilder


def create_sample_deck(theme: str = "dark") -> SlideDeck:
    return SlideDeck(
        deck_title="Executive Growth & Risk Briefing",
        deck_subtitle="Synthesized Data Deck",
        theme=theme,
        author="Enterprise Analytics",
        slides=[
            Slide(
                slide_number=1,
                layout="title_slide",
                title="Executive Growth & Risk Briefing",
                subtitle="Quarterly Overview",
            ),
            Slide(
                slide_number=2,
                layout="kpi_grid",
                title="Key Metrics Scorecard",
                kpi_cards=[
                    KPICard(label="Total ARR", value="$12.4M", change="+14% YoY", trend="up"),
                    KPICard(label="Net Retention", value="108%", change="+2% QoQ", trend="up"),
                ],
                bullet_points=["Strong enterprise momentum across all geos."],
            ),
            Slide(
                slide_number=3,
                layout="chart_and_bullets",
                title="Monthly Performance vs Target",
                chart=ChartConfig(
                    chart_type="bar",
                    categories=["Q1", "Q2", "Q3"],
                    series=[
                        ChartSeries(name="Actual", values=[100, 120, 150]),
                        ChartSeries(name="Target", values=[95, 110, 140]),
                    ],
                ),
                bullet_points=["Consistently beat revenue targets."],
            ),
        ],
    )


def test_pptx_builder_creates_valid_deck():
    deck = create_sample_deck()
    builder = PPTXBuilder(theme_name="dark")
    stream = builder.build_deck(deck)
    assert stream.getbuffer().nbytes > 0

    prs = Presentation(stream)
    assert len(prs.slides) == 3


def test_export_pptx_from_deck_endpoint():
    client = TestClient(app)
    deck = create_sample_deck()
    payload = {
        "deck": deck.model_dump(),
        "filename": "quarterly_briefing",
    }
    response = client.post("/api/v1/presentation/export/pptx", json=payload)
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert "attachment; filename=" in response.headers["content-disposition"]

    prs = Presentation(io.BytesIO(response.content))
    assert len(prs.slides) == 3


def test_export_pptx_from_base64_endpoint():
    client = TestClient(app)
    deck = create_sample_deck()
    builder = PPTXBuilder(theme_name="dark")
    raw_bytes = builder.build_deck(deck).read()
    b64_str = base64.b64encode(raw_bytes).decode("utf-8")

    payload = {
        "base64_data": b64_str,
        "filename": "downloaded_deck.pptx",
    }
    response = client.post("/api/v1/presentation/export/base64", json=payload)
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert len(response.content) == len(raw_bytes)


def test_export_pptx_with_llm_variances_endpoint():
    client = TestClient(app)
    # Payload matching real LLM JSON output variations
    llm_payload = {
        "deck": {
            "title": "Q3 Spend & Risk Briefing",
            "theme": "dark",
            "slides": [
                {
                    "title": "Procurement Highlights",
                    "layout": "bullets",
                    "bullet_points": ["Consolidated top 5 vendor contracts", "Reduced Maverick spend"]
                },
                {
                    "title": "Executive KPIs",
                    "layout": "kpi_grid",
                    "kpis": [
                        {"label": "Total Spend", "val": ".2M", "delta": "-5.2% YoY"},
                        {"label": "Active Vendors", "val": "142", "trend": "down"}
                    ]
                },
                {
                    "title": "Regional Breakdown",
                    "layout": "two_column",
                    "left_column": {
                        "title": "North America",
                        "content": "Spend: .4M<br>Contracts: 85"
                    },
                    "right_column": {
                        "title": "EMEA",
                        "content": "Spend: .8M<br>Contracts: 57"
                    }
                }
            ]
        },
        "filename": "q3_briefing"
    }
    response = client.post("/api/v1/presentation/export/pptx", json=llm_payload)
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert "attachment; filename=" in response.headers["content-disposition"]

    prs = Presentation(io.BytesIO(response.content))
    assert len(prs.slides) >= 3
